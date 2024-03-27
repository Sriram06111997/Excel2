import openpyxl
from openpyxl.chart import (
    LineChart,
    PieChart,
    BarChart,
    Reference
)

wb = openpyxl.Workbook()
ws = wb.active

data = [
    ['Category', 'Value'],
    ['User 1', 1000],
    ['User 2', 2000],
    ['User 3', 3000],
    ['User 4', 4000],
    ['User 5', 5000],
    ['User 6', 6000],
    ['User 7', 7000],
    ['User 8', 8000],
    
]

for row in data:
    ws.append(row)

for _ in range(3):
    ws.append([])

pie_chart = PieChart()
pie_chart.title = "Pie Chart"
labels = Reference(ws, min_col=1, min_row=2, max_row=len(data))
data = Reference(ws, min_col=2, min_row=1, max_row=len(data))
pie_chart.add_data(data, titles_from_data=True)
pie_chart.set_categories(labels)
ws.add_chart(pie_chart, "E1")

for _ in range(3):
    ws.append([])

bar_chart = BarChart()
bar_chart.title = "Bar Chart"
bar_chart.y_axis.title = 'Values'
bar_data = Reference(ws, min_col=2, min_row=1, max_row=len(data))
bar_categories = Reference(ws, min_col=1, min_row=2, max_row=len(data))
bar_chart.add_data(bar_data, titles_from_data=True)
bar_chart.set_categories(bar_categories)
ws.add_chart(bar_chart, "E20")

wb.save("charts.xlsx")





for _ in range(3):
    ws.append([])  


line_chart = LineChart()
line_chart.title = "Line Chart"
line_chart.y_axis.title = 'Values'
line_data = Reference(ws, min_col=2, min_row=1, max_row=len(data))
line_categories = Reference(ws, min_col=1, min_row=2, max_row=len(data))
line_chart.add_data(line_data, titles_from_data=True)
line_chart.set_categories(line_categories)
ws.add_chart(line_chart, "E40")

# Save the workbook
wb.save("charts.xlsx")
